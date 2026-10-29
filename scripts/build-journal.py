"""Build the Word learning journal from the versioned Markdown notes.

Use the bundled document Python runtime resolved by load_workspace_dependencies.
Parent-folder placement is a separate user-approved operation.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parent.parent
document = Document()
section = document.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.85)
normal = document.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string('183B32')
normal.paragraph_format.space_after = Pt(10)
normal.paragraph_format.line_spacing = 1.12
for name in ['Title', 'Heading 1', 'Heading 2']:
    style = document.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
document.core_properties.title = 'GlobeTrek Learning Journal and Interview Practice'
document.core_properties.subject = 'MERN implementation concepts and project interview preparation'
document.add_heading('GlobeTrek Learning Journal', 0)
document.add_paragraph('MERN implementation notes and interview practice', style='Subtitle')
document.add_paragraph('Prepared for the GlobeTrek guided build on 28 October 2026. Use these notes alongside the code and tests, then explain each critical decision in your own words. Hotels and payments in this application are simulated.')
first_topic = True
for source in ['LEARNING.md', 'INTERVIEW.md']:
    content = (root / 'docs' / source).read_text(encoding='utf-8')
    for block in re.split(r'\n\s*\n', content):
        text = block.strip()
        if not text:
            continue
        if text.startswith('# '):
            if source == 'INTERVIEW.md':
                document.add_page_break()
                document.add_heading('GlobeTrek Interview Practice', 0)
            continue
        if text.startswith('## '):
            if not first_topic:
                document.add_page_break()
            heading = re.sub(r'[^\w\s]', '', text[3:])
            document.add_heading(heading, level=1)
            first_topic = False
        else:
            for line in re.split(r'\n(?=\d+\.)', text):
                cleaned = re.sub(r'\s*\n\s*', ' ', line).replace('**', '').replace('`', '')
                paragraph = document.add_paragraph(cleaned)
                paragraph.paragraph_format.widow_control = True
                if re.match(r'^\d+\.', line):
                    paragraph.paragraph_format.keep_together = True
document.save(root / 'docs' / 'GlobeTrek-Learning-Journal.docx')
print('Created docs/GlobeTrek-Learning-Journal.docx')
