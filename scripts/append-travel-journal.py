"""Append the travel expansion without rebuilding or discarding the existing journal."""
from pathlib import Path
import sys
import hashlib
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
root = Path(__file__).resolve().parent.parent
path = root / 'docs' / 'GlobeTrek-Learning-Journal.docx'
parent = root.parent / 'GlobeTrek-Learning-Journal.docx'
source = parent if '--from-parent' in sys.argv and parent.exists() else path
document = Document(source)
# Remove inherited Word title rules, preserving the user's original content.
for element in [document.styles['Title']._element, *[p._p for p in document.paragraphs if p.style.name == 'Title']]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for p in document.paragraphs:
    if p.style.name=='Title':
        for run in p.runs:
            run.font.color.rgb=RGBColor(0,0,0)
            run.font.underline=False
intro='The hotel foundation sections record the earlier learning build. The travel planner sections later in this journal supersede its old 14-night, 90-day and INR-only limits and earlier frontend file locations. Current trips use 1 to 30 days within a 365-day horizon, with explicit destination currencies and server-priced multi-item reservations.'
if not any(p.text==intro for p in document.paragraphs):
    first_topic=next((p for p in document.paragraphs if p.style.name=='Heading 1'),None)
    if first_topic is not None:first_topic.insert_paragraph_before(intro)
if source == parent:
    (root / '.cache').mkdir(exist_ok=True)
    (root / '.cache' / 'journal-parent.sha256').write_text(hashlib.sha256(parent.read_bytes()).hexdigest(), encoding='ascii')
for p in document.paragraphs:
    for run in p.runs:
        corrected=run.text.replace('6. hat belongs in state', '6. What belongs in state').replace('7. race can happen with network responses', '7. What race can happen with network responses')
        if corrected != run.text:run.text=corrected
for filename, marker in [('TRAVEL-LEARNING.md', 'GlobeTrek travel planner learning notes'), ('RELEASE-LEARNING.md', 'Release hardening and recovery'), ('CONNECTED-MAPS-LEARNING.md', 'Connected maps and place previews'), ('SCROLLING-LEARNING.md', 'Momentum scrolling')]:
  if not any(p.text == marker for p in document.paragraphs):
    document.add_page_break()
    for line in (root / 'docs' / filename).read_text(encoding='utf-8').splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            document.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            document.add_heading(line[3:], level=2)
        else:
            paragraph = document.add_paragraph(line.replace('**', '').replace('`', ''))
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.widow_control = True
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
document.save(path)
print(path)
