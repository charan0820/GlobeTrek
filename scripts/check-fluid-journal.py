import io,json,subprocess
from pathlib import Path
from docx import Document
previous=Document(io.BytesIO(subprocess.check_output(['git','show','HEAD:docs/GlobeTrek-Learning-Journal.docx'])))
current=Document('docs/GlobeTrek-Learning-Journal.docx')
a=[p.text for p in previous.paragraphs if p.text.strip()];b=[p.text for p in current.paragraphs if p.text.strip()]
r={'preservedParagraphPrefix':b[:len(a)]==a,'previousParagraphs':len(a),'currentParagraphs':len(b),'visualVerification':'Blocked: LibreOffice unavailable; hidden Word renderer stalled and was cancelled.'}
assert r['preservedParagraphPrefix']
Path('docs/evidence/fluid-atlas/journal-qa.json').write_text(json.dumps(r,indent=2),encoding='utf-8');print(r)
