from pathlib import Path
from docx import Document
from docx.shared import Pt
path=Path('docs/GlobeTrek-Learning-Journal.docx')
document=Document(path)
marker='Fluid atlas layouts and contextual islands'
if not any(p.text==marker for p in document.paragraphs):
 document.add_page_break()
 for line in Path('docs/FLUID-ATLAS-LEARNING.md').read_text(encoding='utf-8').splitlines():
  if line.startswith('# '): document.add_heading(line[2:],level=1)
  elif line.startswith('## '): document.add_heading(line[3:],level=2)
  elif line.strip():
   p=document.add_paragraph(line)
   p.paragraph_format.space_after=Pt(8)
   p.paragraph_format.widow_control=True
 document.save(path)
print(path)
